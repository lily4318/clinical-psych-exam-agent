"""Regression tests for all-or-nothing JSON-to-two-DOCX publication.

Synthetic questions only; these checks do not validate clinical content.
"""

import os

import pytest
from docx import Document

from src import export_docx


def _exam(form):
    return {
        "form": form,
        "version": "atomic-test-v1",
        "approved_by_user": False,
        "items": [{
            "pair_id": "P01",
            "stem": f"{form} 合成測試題幹",
            "options": {letter: f"{letter} 測試選項" for letter in "ABCD"},
            "correct_answer": "A",
            "rationale": {
                "correct": "這是測試用的正解說明。",
                **{letter: f"{letter} 選項的測試解析" for letter in "ABCD"},
            },
            "sources": ["UNIT-TEST-NOT-CLINICAL"],
        }],
    }


def _assert_no_staging(tmp_path):
    assert not list(tmp_path.glob(".parallel-docx-stage-*"))


def test_parallel_export_commits_two_openable_documents_without_staging(tmp_path):
    paths = export_docx.export_parallel_docx(_exam("甲卷"), _exam("乙卷"), tmp_path)
    assert set(paths) == {"form_a", "form_b"}
    assert len(list(tmp_path.glob("*.docx"))) == 2
    assert "甲卷" in "\n".join(p.text for p in Document(paths["form_a"]).paragraphs)
    assert "乙卷" in "\n".join(p.text for p in Document(paths["form_b"]).paragraphs)
    _assert_no_staging(tmp_path)


def test_failure_while_staging_second_word_leaves_no_first_word(tmp_path, monkeypatch):
    original = export_docx.export_exam_docx

    def fail_on_b(exam, out_dir, **kwargs):
        if exam["form"] == "乙卷":
            raise OSError("simulated second staging failure")
        return original(exam, out_dir, **kwargs)

    monkeypatch.setattr(export_docx, "export_exam_docx", fail_on_b)
    with pytest.raises(OSError, match="second staging failure"):
        export_docx.export_parallel_docx(_exam("甲卷"), _exam("乙卷"), tmp_path)
    assert not list(tmp_path.glob("*.docx"))
    _assert_no_staging(tmp_path)


def test_failure_publishing_second_word_rolls_back_first(tmp_path, monkeypatch):
    original = os.link

    def fail_on_b(source, destination, *args, **kwargs):
        if destination.name.startswith("form_b_"):
            raise OSError("simulated second publication failure")
        return original(source, destination, *args, **kwargs)

    monkeypatch.setattr(export_docx.os, "link", fail_on_b)
    with pytest.raises(OSError, match="second publication failure"):
        export_docx.export_parallel_docx(_exam("甲卷"), _exam("乙卷"), tmp_path)
    assert not list(tmp_path.glob("*.docx"))
    _assert_no_staging(tmp_path)


def test_existing_word_is_not_overwritten_or_paired_with_new_file(tmp_path):
    original = tmp_path / "form_a_atomic-test-v1.docx"
    original.write_bytes(b"existing-owner-data")
    with pytest.raises(FileExistsError, match="Refusing to overwrite"):
        export_docx.export_parallel_docx(_exam("甲卷"), _exam("乙卷"), tmp_path)
    assert original.read_bytes() == b"existing-owner-data"
    assert not (tmp_path / "form_b_atomic-test-v1.docx").exists()
    _assert_no_staging(tmp_path)


def test_repeating_export_refuses_to_replace_either_delivered_word(tmp_path):
    paths = export_docx.export_parallel_docx(_exam("甲卷"), _exam("乙卷"), tmp_path)
    before = {key: path.read_bytes() for key, path in paths.items()}
    with pytest.raises(FileExistsError, match="Refusing to overwrite"):
        export_docx.export_parallel_docx(_exam("甲卷"), _exam("乙卷"), tmp_path)
    assert all(paths[key].read_bytes() == data for key, data in before.items())
    assert len(list(tmp_path.glob("*.docx"))) == 2
    _assert_no_staging(tmp_path)
