"""Guard against duplicated A-D option content in native and external DOCX.

Synthetic fixtures only; no clinical claims, approved exams or paid inference.
"""

import pytest
from docx import Document

from src.export_docx import export_parallel_docx
from src.verify_external_delivery import verify_pair


def _form(label):
    return {
        "form": label,
        "version": "synthetic-distinct-options",
        "domain_id": "child_adolescent",
        "approved_by_user": False,
        "items": [{
            "pair_id": "P01",
            "stem": f"{label} 合成題幹，非臨床試題",
            "options": {key: f"合成選項 {key}" for key in "ABCD"},
            "correct_answer": "C",
            "rationale": {
                "correct": "僅供程式驗證的合成說明。",
                **{key: f"合成選項 {key} 的個別解析。" for key in "ABCD"},
            },
            "sources": ["SYNTHETIC-TEST-SOURCE"],
        }],
    }


@pytest.mark.parametrize(
    ("option_a", "option_b"),
    [
        ("相同測試內容", "相同測試內容"),
        ("Ａ型測試內容", "A型測試內容"),
        ("測 試 內容", "測試內容"),
    ],
)
@pytest.mark.parametrize("include_answers", [True, False])
def test_native_export_refuses_duplicate_option_content_before_writes(
    tmp_path, option_a, option_b, include_answers
):
    a, b = _form("甲卷"), _form("乙卷")
    b["items"][0]["options"]["A"] = option_a
    b["items"][0]["options"]["B"] = option_b
    with pytest.raises(ValueError, match="four distinct A-D option texts"):
        export_parallel_docx(a, b, tmp_path, include_answers=include_answers)
    assert not list(tmp_path.glob("*.docx"))


@pytest.mark.parametrize("variant", ["exact", "whitespace"])
def test_external_verifier_refuses_duplicate_option_text_after_docx_edit(
    tmp_path, variant
):
    paths = export_parallel_docx(_form("甲卷"), _form("乙卷"), tmp_path)
    document = Document(paths["form_b"])
    option_a = next(
        paragraph.text.removeprefix("(A) ")
        for paragraph in document.paragraphs
        if paragraph.text.startswith("(A) ")
    )
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("(B) "):
            paragraph.text = "(B) " + (
                option_a if variant == "exact" else "  " + option_a.replace(" ", "　") + "  "
            )
            break
    else:
        pytest.fail("Expected option B paragraph in the synthetic Word fixture.")
    document.save(paths["form_b"])

    report = verify_pair(paths["form_a"], paths["form_b"], 1)
    assert not report["ok"]
    assert any(
        "A–D 選項內容重複" in error
        for error in report["files"]["form_b"]["errors"]
    )
