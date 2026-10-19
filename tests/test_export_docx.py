import pytest
from docx import Document

from src.export_docx import export_exam_docx, export_parallel_docx


def _exam(form: str) -> dict:
    item = {
        "pair_id": "P01",
        "stem": "兒童心理衡鑑應整合哪些資訊？",
        "options": {
            "A": "僅家長評估",
            "B": "只有學校觀察",
            "C": "多資訊來源與發展史",
            "D": "只看智力分數",
        },
        "correct_answer": "C",
        "rationale": {
            "correct": "應整合不同來源與發展史。",
            "A": "資訊不足。",
            "B": "資訊不足。",
            "C": "符合完整衡鑑原則。",
            "D": "智力分數不等於整體功能。",
        },
        "sources": ["TEST-SRC-001"],
    }
    return {
        "form": form,
        "version": "child-review-v0.1",
        "items": [item],
        "approved_by_user": False,
    }


def test_exports_two_real_word_documents_with_answers(tmp_path):
    paths = export_parallel_docx(_exam("甲卷"), _exam("乙卷"), tmp_path)
    assert set(paths) == {"form_a", "form_b"}
    assert len(set(paths.values())) == 2

    for form, path in (("甲卷", paths["form_a"]), ("乙卷", paths["form_b"])):
        assert path.exists() and path.suffix == ".docx"
        doc = Document(path)
        paragraph_text = "\n".join(p.text for p in doc.paragraphs)
        assert f"臨床心理師模擬測驗｜{form}" in paragraph_text
        assert "兒童心理衡鑑應整合哪些資訊？" in paragraph_text
        assert "(C) 多資訊來源與發展史" in paragraph_text
        assert f"{form}｜答案與解析" in paragraph_text
        assert "正解解析：應整合不同來源與發展史。" in paragraph_text
        assert "來源 ID：TEST-SRC-001" in paragraph_text
        assert len(doc.tables) == 1
        assert doc.tables[0].rows[1].cells[1].text == "C"
        assert "未標示為使用者正式核可版本" in paragraph_text


def test_exam_only_word_hides_answer_key_and_rationale(tmp_path):
    path = export_exam_docx(_exam("甲卷"), tmp_path, include_answers=False)
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "兒童心理衡鑑應整合哪些資訊？" in text
    assert "答案與解析" not in text
    assert "正解解析" not in text
    assert "來源 ID" not in text
    assert not doc.tables


def test_rejects_mismatched_parallel_pairs(tmp_path):
    form_a, form_b = _exam("甲卷"), _exam("乙卷")
    form_b["items"][0]["pair_id"] = "P02"
    with pytest.raises(ValueError, match="Pair ID mismatch"):
        export_parallel_docx(form_a, form_b, tmp_path)
    assert not list(tmp_path.glob("*.docx"))


def test_rejects_invalid_answers_before_writing(tmp_path):
    form_a, form_b = _exam("甲卷"), _exam("乙卷")
    form_b["items"][0]["correct_answer"] = "E"
    with pytest.raises(ValueError, match="valid answer letter"):
        export_parallel_docx(form_a, form_b, tmp_path)
    assert not list(tmp_path.glob("*.docx"))


def test_existing_review_manifest_exports_complete_two_word_files(tmp_path):
    # End-to-end from the frozen review manifest; no model/API calls.
    from src.export_review_docx import export_review_docx

    paths = export_review_docx(tmp_path)
    assert set(paths) == {"form_a", "form_b"}
    assert len(list(tmp_path.glob("*.docx"))) == 2

    for label, expected_form in (("form_a", "甲卷"), ("form_b", "乙卷")):
        doc = Document(paths[label])
        texts = [paragraph.text for paragraph in doc.paragraphs]
        assert any(f"臨床心理師模擬測驗｜{expected_form}" in t for t in texts)
        assert any(f"{expected_form}｜答案與解析" in t for t in texts)
        assert any("未標示為使用者正式核可版本" in t for t in texts)
        assert sum(t.startswith("正解解析：") for t in texts) == 20
        assert len(doc.tables) == 1
        assert len(doc.tables[0].rows) == 21
        assert len(paths[label].read_bytes()) > 5000
