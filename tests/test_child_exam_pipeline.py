"""Purely mocked inference: exercises the child draft path, never clinical validity."""

from copy import deepcopy

import pytest
from docx import Document

from src.child_exam_pipeline import generate_child_parallel_docx
from src.domain_blueprints import load_domain_blueprint


def _evidence():
    blueprint = load_domain_blueprint()
    return {
        sid: {
            "url": source["url"],
            "retrieved_at": "2026-10-20",
            "passage": (
                "這是單元測試專用的合成來源摘錄，僅測試程式的資料完整性和來源身分。"
                "內容不構成臨床事實、診斷準則或治療建議。"
            ),
        }
        for sid, source in blueprint["sources"].items()
        if source["kind"] == "clinical_guideline"
    }


def _item(form, spec, slot, source_id):
    return {
        "item_id": slot["item_id"],
        "form": form,
        "pair_id": spec["pair_id"],
        "stem": f"{form} 程式回歸測試情境，配對 {spec['pair_id']}；非真實臨床試題。",
        "options": {
            "A": "測試選項一",
            "B": "測試選項二",
            "C": "測試選項三",
            "D": "測試選項四",
        },
        "correct_answer": slot["target_answer_position"],
        "category": spec["category"],
        "subtopic": spec["core_topic"],
        "difficulty": spec["difficulty"],
        "cognitive_level": spec["cognitive_level"],
        "rationale": {
            "correct": "合成測試中指定的位置是預期答案，不代表臨床事實。",
            "A": "測試A選項說明",
            "B": "測試B選項說明",
            "C": "測試C選項說明",
            "D": "測試D選項說明",
        },
        "sources": [source_id],
        "quality_status": "草稿",
    }


def _valid_response(request):
    spec = request["pair_spec"]
    source_id = request["source_context"][0]["source_id"]
    return {
        "pair_id": spec["pair_id"],
        "form_a": _item("甲卷", spec, request["forms"][0], source_id),
        "form_b": _item("乙卷", spec, request["forms"][1], source_id),
    }


class FakeModel:
    def __init__(self, mutator=None):
        self.calls = []
        self.mutator = mutator

    def __call__(self, request):
        self.calls.append(request)
        result = _valid_response(request)
        if self.mutator:
            self.mutator(result, request)
        return result


@pytest.mark.parametrize("count", [1, 7, 20, 23])
def test_child_pipeline_delivers_two_real_word_files_with_answers(tmp_path, count):
    model = FakeModel()
    result = generate_child_parallel_docx(
        model, _evidence(), tmp_path, items_per_form=count
    )
    assert len(model.calls) == count
    assert set(result) == {"form_a", "form_b"}
    assert len(list(tmp_path.glob("*.docx"))) == 2
    assert len(set(result.values())) == 2
    for key, form in (("form_a", "甲卷"), ("form_b", "乙卷")):
        path = result[key]
        assert path.exists() and path.suffix == ".docx"
        doc = Document(path)
        paragraphs = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        assert form in paragraphs
        assert "答案與解析" in paragraphs
        assert "命題範圍與依據" in paragraphs
        assert "命題範圍：兒童／青少年" in paragraphs
        assert "考點配置：" in paragraphs
        assert "來源對照與參考資料" in paragraphs
        assert "第 1 題：" in paragraphs
        assert "未經獨立臨床專家逐題驗證" in paragraphs
        assert "版本：" not in paragraphs
        assert "未標示為使用者正式核可版本" not in paragraphs
        assert paragraphs.count("正解解析：") == count
        assert len(doc.tables[0].rows) == count + 1

    first = model.calls[0]
    assert first["pair_spec"]["category"] == "ADHD"
    assert first["forms"][0]["target_answer_position"] == "A"
    assert first["forms"][1]["target_answer_position"] == "C"
    assert all(
        ctx["verification_status"] == "provided_excerpt_not_semantically_verified"
        for call in model.calls for ctx in call["source_context"]
    )


def test_source_missing_fails_before_any_inference_or_output(tmp_path):
    model = FakeModel()
    evidence = _evidence()
    evidence.pop("NICE-NG87")
    with pytest.raises(ValueError, match="missing clinical source passage"):
        generate_child_parallel_docx(model, evidence, tmp_path)
    assert model.calls == []
    assert not list(tmp_path.glob("*.docx"))


@pytest.mark.parametrize("bad_evidence", [
    {"url": "https://example.com/unrelated", "passage": "夠長的原文內容" * 12, "retrieved_at": "2026-10-20"},
    {"url": None, "passage": "夠長的原文內容" * 12, "retrieved_at": "2026-10-20"},
    {"url": "MATCH", "passage": "too short", "retrieved_at": "2026-10-20"},
    {"url": "MATCH", "passage": "[STUB] 替代文字" * 8, "retrieved_at": "2026-10-20"},
    {"url": "MATCH", "passage": "資料來源原文引用測試" * 8, "retrieved_at": ""},
])
def test_inadequate_or_mismatched_evidence_fails_closed(tmp_path, bad_evidence):
    evidence = _evidence()
    modified = deepcopy(bad_evidence)
    if modified.get("url") == "MATCH":
        modified["url"] = evidence["NICE-NG87"]["url"]
    evidence["NICE-NG87"] = modified
    model = FakeModel()
    with pytest.raises(ValueError):
        generate_child_parallel_docx(model, evidence, tmp_path, items_per_form=1)
    assert model.calls == []
    assert not list(tmp_path.glob("*.docx"))


@pytest.mark.parametrize("value", [0, -1, True, "10"])
def test_invalid_question_count_does_not_generate(tmp_path, value):
    model = FakeModel()
    with pytest.raises(ValueError, match="positive integer"):
        generate_child_parallel_docx(model, _evidence(), tmp_path, items_per_form=value)
    assert not model.calls
    assert not list(tmp_path.glob("*.docx"))


@pytest.mark.parametrize("version", ["", "formal-v1.0", "child-approved", "核可草稿"])
def test_formal_version_spoofing_disallowed(tmp_path, version):
    model = FakeModel()
    with pytest.raises(ValueError, match="Draft|labelled"):
        generate_child_parallel_docx(model, _evidence(), tmp_path, version=version)
    assert not model.calls
    assert not list(tmp_path.glob("*.docx"))


def _mutate(target, value):
    def callback(result, request):
        result["form_b"][target] = value
    return callback


@pytest.mark.parametrize("mutator,reason", [
    (_mutate("correct_answer", "E"), "correct_answer"),
    (_mutate("correct_answer", "A"), "correct_answer"),
    (_mutate("quality_status", "已核可"), "quality_status"),
    (_mutate("category", "C"), "category"),
    (_mutate("sources", ["UNKNOWN-SOURCE"]), "Sources"),
    (_mutate("sources", []), "Sources"),
    (_mutate("rationale", {"correct": "只有一個說明"}), "Rationale"),
    (_mutate("options", {"A": "同", "B": "同", "C": "異", "D": "其他"}), "Options"),
    (_mutate("stem", "[STUB] 題目"), "stem"),
    (_mutate("approved_by_user", True), "formal approval"),
])
def test_invalid_model_output_never_writes_word(tmp_path, mutator, reason):
    model = FakeModel(mutator)
    with pytest.raises(ValueError, match=reason):
        generate_child_parallel_docx(model, _evidence(), tmp_path, items_per_form=1)
    assert len(model.calls) == 1
    assert not list(tmp_path.glob("*.docx"))


def test_model_pair_id_mismatch_blocks_export(tmp_path):
    model = FakeModel(lambda response, request: response.update(pair_id="P88"))
    with pytest.raises(ValueError, match="pair_id mismatch"):
        generate_child_parallel_docx(model, _evidence(), tmp_path, items_per_form=1)
    assert not list(tmp_path.glob("*.docx"))


def test_parallel_stems_must_not_be_identical(tmp_path):
    model = FakeModel(
        lambda response, request: response["form_b"].update(
            stem=response["form_a"]["stem"]
        )
    )
    with pytest.raises(ValueError, match="parallel question stems are identical"):
        generate_child_parallel_docx(model, _evidence(), tmp_path, items_per_form=1)
    assert not list(tmp_path.glob("*.docx"))


def test_source_insufficient_halts_before_writing(tmp_path):
    def mark_insufficient(response, request):
        response["source_insufficient"] = True
    model = FakeModel(mark_insufficient)
    with pytest.raises(ValueError, match="could not supply a valid pair"):
        generate_child_parallel_docx(model, _evidence(), tmp_path, items_per_form=1)
    assert not list(tmp_path.glob("*.docx"))


def test_late_failure_does_not_leave_first_word_file(tmp_path):
    def fail_second_pair(response, request):
        if request["pair_spec"]["pair_id"] == "P02":
            response["form_a"]["sources"] = ["UNRELATED"]
    model = FakeModel(fail_second_pair)
    with pytest.raises(ValueError, match="Sources"):
        generate_child_parallel_docx(model, _evidence(), tmp_path, items_per_form=2)
    assert len(model.calls) == 2
    assert not list(tmp_path.glob("*.docx"))


def test_cannot_fall_back_to_stub_when_no_callable(tmp_path):
    with pytest.raises(TypeError, match="actual model callback"):
        generate_child_parallel_docx(None, _evidence(), tmp_path, items_per_form=1)
    assert not list(tmp_path.glob("*.docx"))


def test_default_20q_full_pipeline_strictly_verifies_both_word_files(tmp_path):
    """A 20-item structure smoke with FakeModel, NOT a live clinical inference trial."""
    from collections import Counter

    from src.domain_blueprints import build_child_pair_specs
    from src.verify_external_delivery import verify_pair

    model = FakeModel()
    paths = generate_child_parallel_docx(model, _evidence(), tmp_path)
    assert len(model.calls) == 20
    assert len(list(tmp_path.glob("*.docx"))) == 2
    assert not list(tmp_path.glob(".child-docx-stage-*"))

    report = verify_pair(paths["form_a"], paths["form_b"], 20)
    assert report["ok"], report
    for key, form in (("form_a", "甲卷"), ("form_b", "乙卷")):
        entry = report["files"][key]
        assert entry["reader_layout"] == "modern_strict"
        assert entry["questions"] == 20
        assert entry["answers"] == 20
        assert entry["complete_option_rationales"] == 20
        assert not entry["errors"]
        doc = Document(paths[key])
        paragraph_text = "\n".join(p.text for p in doc.paragraphs)
        assert doc.core_properties.identifier.startswith("exam_version:")
        assert f"{form}｜答案與解析" in paragraph_text
        assert paragraph_text.count("正解解析：") == 20
        assert paragraph_text.count("A 選項：") == 20
        assert paragraph_text.count("B 選項：") == 20
        assert paragraph_text.count("C 選項：") == 20
        assert paragraph_text.count("D 選項：") == 20
        assert len(doc.tables[0].rows) == 21
        assert "未經獨立臨床專家逐題驗證" in paragraph_text
        assert "版本：" not in paragraph_text

    planned = build_child_pair_specs(20)
    actual = Counter(call["pair_spec"]["category"] for call in model.calls)
    expected = Counter(spec["category"] for spec in planned)
    assert actual == expected
    assert sum(actual.values()) == 20
    for request in model.calls:
        assert request["source_context"]
        assert all(
            source["verification_status"] == "provided_excerpt_not_semantically_verified"
            for source in request["source_context"]
        )


def test_post_export_structural_failure_publishes_neither_word(tmp_path, monkeypatch):
    """A verifier failure after 20 synthetic pairs cannot become a delivered paper."""
    from src import child_exam_pipeline

    def fail_structural_check(*args, **kwargs):
        return {
            "ok": False,
            "files": {
                "form_a": {"ok": True, "errors": []},
                "form_b": {"ok": False, "errors": ["missing explanation"]},
            },
        }

    monkeypatch.setattr(child_exam_pipeline, "verify_pair", fail_structural_check)
    model = FakeModel()
    with pytest.raises(ValueError, match="strict structural verification"):
        generate_child_parallel_docx(model, _evidence(), tmp_path)
    assert len(model.calls) == 20
    assert not list(tmp_path.glob("*.docx"))
    assert not list(tmp_path.glob(".child-docx-stage-*"))


def test_second_publication_failure_removes_first_word(tmp_path, monkeypatch):
    from pathlib import Path
    from src import child_exam_pipeline

    original = child_exam_pipeline.os.link

    def block_second_file(source, target):
        if Path(target).parent == tmp_path and Path(target).name.startswith("form_b_"):
            raise OSError("simulated second-file publication failure")
        return original(source, target)

    monkeypatch.setattr(child_exam_pipeline.os, "link", block_second_file)
    with pytest.raises(OSError, match="second-file publication failure"):
        generate_child_parallel_docx(
            FakeModel(), _evidence(), tmp_path, items_per_form=2
        )
    assert not list(tmp_path.glob("*.docx"))
    assert not list(tmp_path.glob(".child-docx-stage-*"))


@pytest.mark.parametrize("blocked_form", ["a", "b"])
def test_competing_publisher_cannot_be_overwritten(tmp_path, monkeypatch, blocked_form):
    """The final path may appear after exists() but before publication."""
    from pathlib import Path
    from src import child_exam_pipeline

    original = child_exam_pipeline.os.link
    injected = []

    def race_at_link(source, target):
        destination = Path(target)
        if destination.parent == tmp_path and destination.name.startswith(f"form_{blocked_form}_"):
            destination.write_bytes(b"another publisher's document")
            injected.append(destination)
            raise FileExistsError("simulated competing publisher")
        return original(source, target)

    monkeypatch.setattr(child_exam_pipeline.os, "link", race_at_link)
    with pytest.raises(FileExistsError, match="competing publisher"):
        generate_child_parallel_docx(
            FakeModel(), _evidence(), tmp_path, items_per_form=2
        )
    assert len(injected) == 1
    assert injected[0].read_bytes() == b"another publisher's document"
    assert list(tmp_path.glob("*.docx")) == injected
    assert not list(tmp_path.glob(".child-docx-stage-*"))


def test_unsupported_hard_links_fail_without_partial_delivery(tmp_path, monkeypatch):
    from src import child_exam_pipeline

    def unsupported(*args, **kwargs):
        raise OSError("hard links unsupported")

    monkeypatch.setattr(child_exam_pipeline.os, "link", unsupported)
    with pytest.raises(OSError, match="hard links unsupported"):
        generate_child_parallel_docx(
            FakeModel(), _evidence(), tmp_path, items_per_form=2
        )
    assert not list(tmp_path.glob("*.docx"))
    assert not list(tmp_path.glob(".child-docx-stage-*"))


def test_same_version_cannot_overwrite_existing_delivered_paper(tmp_path):
    first = generate_child_parallel_docx(
        FakeModel(), _evidence(), tmp_path, items_per_form=1
    )
    previous = {key: path.read_bytes() for key, path in first.items()}
    with pytest.raises(FileExistsError, match="Refusing to overwrite"):
        generate_child_parallel_docx(
            FakeModel(), _evidence(), tmp_path, items_per_form=1
        )
    assert len(list(tmp_path.glob("*.docx"))) == 2
    assert all(first[key].read_bytes() == value for key, value in previous.items())
    assert not list(tmp_path.glob(".child-docx-stage-*"))


@pytest.mark.parametrize("reused_form", ["form_a", "form_b"])
def test_repeated_stem_across_different_pairs_blocks_both_word_files(
    tmp_path, reused_form
):
    """Model can repeat a previous stem even when A/B of each pair differ."""
    first_stem = {}

    def repeat_previous(response, request):
        if request["pair_spec"]["pair_id"] == "P01":
            first_stem["value"] = response["form_a"]["stem"]
        else:
            response[reused_form]["stem"] = "  " + first_stem["value"] + "　"

    model = FakeModel(repeat_previous)
    with pytest.raises(ValueError, match="duplicate question stem"):
        generate_child_parallel_docx(model, _evidence(), tmp_path, items_per_form=2)
    assert len(model.calls) == 2
    assert not list(tmp_path.glob("*.docx"))


def test_same_pair_fullwidth_normalized_stems_are_rejected(tmp_path):
    def repeat_fullwidth(response, request):
        response["form_b"]["stem"] = (
            "　" + response["form_a"]["stem"].replace("P01", "Ｐ０１") + "　"
        )

    with pytest.raises(ValueError, match="parallel question stems are identical"):
        generate_child_parallel_docx(
            FakeModel(repeat_fullwidth), _evidence(), tmp_path, items_per_form=1
        )
    assert not list(tmp_path.glob("*.docx"))
