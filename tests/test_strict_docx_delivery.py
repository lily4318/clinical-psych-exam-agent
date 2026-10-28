"""DOCX completion gates: synthetic formatting fixtures, never released exam questions."""

from copy import deepcopy

import pytest
from docx import Document

from src.export_docx import export_parallel_docx
from src.verify_external_delivery import verify_pair


def _form(form: str, count: int = 2) -> dict:
    return {
        "form": form,
        "version": "strict-reader-test-v1",
        "domain_id": "child_adolescent",
        "domain_name": "兒童／青少年臨床心理",
        "approved_by_user": False,
        "items": [
            {
                "pair_id": f"P{i:02d}",
                "stem": f"{form} 匯出驗證測試題目 {i}，非實際臨床試題",
                "options": {letter: f"選項 {letter} 的測試文字" for letter in "ABCD"},
                "correct_answer": "C",
                "rationale": {
                    "correct": "這是測試用的結構化解析，不代表真實診斷依據。",
                    **{letter: f"{letter} 選項測試用個別解析。" for letter in "ABCD"},
                },
                "category": "TEST",
                "sources": ["TEST-SOURCE-001"],
            }
            for i in range(1, count + 1)
        ],
    }


def _pair(tmp_path):
    return export_parallel_docx(_form("甲卷"), _form("乙卷"), tmp_path)


@pytest.mark.parametrize("letter", list("ABCD"))
def test_export_refuses_missing_option_rationale_before_any_writes(tmp_path, letter):
    a, b = _form("甲卷"), _form("乙卷")
    b["items"][1]["rationale"].pop(letter)
    with pytest.raises(ValueError, match=f"{letter} rationale"):
        export_parallel_docx(a, b, tmp_path)
    assert not list(tmp_path.glob("*.docx"))


@pytest.mark.parametrize("bad", [None, [], [""], ["TEST-SOURCE-001", "TEST-SOURCE-001"]])
def test_export_refuses_invalid_sources_before_any_writes(tmp_path, bad):
    a, b = _form("甲卷"), _form("乙卷")
    if bad is None:
        b["items"][0].pop("sources")
    else:
        b["items"][0]["sources"] = bad
    with pytest.raises(ValueError, match="source IDs"):
        export_parallel_docx(a, b, tmp_path)
    assert not list(tmp_path.glob("*.docx"))


def test_complete_modern_documents_are_structurally_accepted(tmp_path):
    pair = _pair(tmp_path)
    result = verify_pair(pair["form_a"], pair["form_b"], 2)
    assert result["ok"], result
    for form in result["files"].values():
        assert form["reader_layout"] == "modern_strict"
        assert form["complete_option_rationales"] == 2
        assert not form["warnings"]
    assert result["clinical_correctness_and_source_support"] == "not_independently_verified"


def test_modern_missing_option_rationale_fails(tmp_path):
    pair = _pair(tmp_path)
    path = pair["form_b"]
    doc = Document(path)
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("B 選項："):
            paragraph.text = ""
            break
    doc.save(path)
    result = verify_pair(pair["form_a"], path, 2)
    assert not result["ok"]
    assert any("B" in e and "逐項解析" in e for e in result["files"]["form_b"]["errors"])


def test_modern_blank_option_rationale_fails(tmp_path):
    pair = _pair(tmp_path)
    path = pair["form_b"]
    doc = Document(path)
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("D 選項："):
            paragraph.text = "D 選項：  "
            break
    doc.save(path)
    assert not verify_pair(pair["form_a"], path, 2)["ok"]


def test_modern_wrong_section_order_fails(tmp_path):
    pair = _pair(tmp_path)
    path = pair["form_b"]
    doc = Document(path)
    headings = [p for p in doc.paragraphs if p.text in ("命題範圍與依據", "答案一覽")]
    assert len(headings) == 2
    headings[0].text, headings[1].text = headings[1].text, headings[0].text
    doc.save(path)
    result = verify_pair(pair["form_a"], path, 2)
    assert not result["ok"]
    assert any("章節順序" in error for error in result["files"]["form_b"]["errors"])


def test_modern_missing_source_mapping_fails(tmp_path):
    pair = _pair(tmp_path)
    path = pair["form_b"]
    doc = Document(path)
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("第 2 題："):
            paragraph.text = ""
    doc.save(path)
    result = verify_pair(pair["form_a"], path, 2)
    assert not result["ok"]
    assert any("來源" in error for error in result["files"]["form_b"]["errors"])



@pytest.mark.parametrize(
    ("source_line", "expected_error"),
    [
        ("第 1 題：TEST-SOURCE-001, TEST-SOURCE-001", "來源 ID 重複"),
        ("第 1 題：TEST-SOURCE-001, ", "來源 ID 不可為空"),
    ],
)
def test_modern_source_appendix_rejects_duplicate_or_empty_ids(
    tmp_path, source_line, expected_error
):
    """A structurally present source mapping is not sufficient if IDs are invalid."""
    pair = _pair(tmp_path)
    path = pair["form_b"]
    doc = Document(path)
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("第 1 題："):
            paragraph.text = source_line
            break
    else:
        pytest.fail("Fixture did not contain the first source mapping.")
    doc.save(path)
    result = verify_pair(pair["form_a"], path, 2)
    assert not result["ok"]
    assert any(
        expected_error in error
        for error in result["files"]["form_b"]["errors"]
    )


def test_modern_source_appendix_before_explanations_fails(tmp_path):
    pair = _pair(tmp_path)
    path = pair["form_b"]
    doc = Document(path)
    titles = [p for p in doc.paragraphs if p.text in ("逐題解析", "來源對照與參考資料")]
    assert len(titles) == 2
    titles[0].text, titles[1].text = titles[1].text, titles[0].text
    doc.save(path)
    result = verify_pair(pair["form_a"], path, 2)
    assert not result["ok"]
    assert any("章節順序" in error or "逐題解析之後" in error for error in result["files"]["form_b"]["errors"])


def test_modern_preamble_contains_no_ai_generation_metadata(tmp_path):
    pair = _pair(tmp_path)
    path = pair["form_b"]
    doc = Document(path)
    doc.paragraphs[1].text = "版本：內部管線測試"
    doc.save(path)
    result = verify_pair(pair["form_a"], path, 2)
    assert not result["ok"]
    assert any("試題區" in error for error in result["files"]["form_b"]["errors"])


def test_legacy_files_can_still_be_read_but_not_called_strict(tmp_path):
    """Preserve backwards compatibility without claiming full modern acceptance."""
    legacy = {}
    for form, filename in (("甲卷", "a.docx"), ("乙卷", "b.docx")):
        doc = Document()
        doc.add_paragraph(f"臨床心理師模擬測驗｜{form}")
        doc.add_paragraph("版本：old-test-v0")
        doc.add_paragraph("未經獨立專家驗證（測試）")
        doc.add_paragraph(f"1. {form} 歷史版面題幹")
        for letter in "ABCD":
            doc.add_paragraph(f"({letter}) 測試選項 {letter}")
        doc.add_paragraph(f"{form}｜答案與解析")
        tbl = doc.add_table(rows=2, cols=2)
        tbl.rows[0].cells[0].text, tbl.rows[0].cells[1].text = "題號", "答案"
        tbl.rows[1].cells[0].text, tbl.rows[1].cells[1].text = "1", "A"
        doc.add_paragraph("1. A")
        doc.add_paragraph("正解解析：測試用解釋")
        doc.add_paragraph("資料來源 ID：TEST-LEGACY")
        path = tmp_path / filename
        doc.save(path)
        legacy[form] = path
    report = verify_pair(legacy["甲卷"], legacy["乙卷"], 1)
    assert report["ok"], report
    for form in report["files"].values():
        assert form["reader_layout"] == "legacy_compatible"
        assert form["complete_option_rationales"] is None
        assert form["warnings"]


@pytest.mark.parametrize("repeat_target", ["same_form", "other_form"])
def test_external_docx_rejects_reused_stem_at_different_question_number(
    tmp_path, repeat_target
):
    """Externally generated Word cannot silently repeat a question elsewhere."""
    a, b = _form("甲卷"), _form("乙卷")
    repeated = a["items"][0]["stem"]
    if repeat_target == "same_form":
        a["items"][1]["stem"] = "　" + repeated + "  "
    else:
        b["items"][1]["stem"] = repeated.replace("1", "１")
    pair = export_parallel_docx(a, b, tmp_path)
    result = verify_pair(pair["form_a"], pair["form_b"], 2)
    assert not result["ok"]
    form = "form_a" if repeat_target == "same_form" else "form_b"
    assert any("題幹重複" in e for e in result["files"][form]["errors"])
    assert result["files"][form]["questions"] == 2
