"""Exercise the public DOCX command-line entrypoint without model/API calls."""

import json
import subprocess
import sys

from docx import Document


def _exam(form: str) -> dict:
    return {
        "form": form,
        "version": "cli-test-v1",
        "approved_by_user": False,
        "items": [
            {
                "pair_id": "P01",
                "stem": "測試題幹",
                "options": {
                    "A": "選項甲",
                    "B": "選項乙",
                    "C": "選項丙",
                    "D": "選項丁",
                },
                "correct_answer": "C",
                "rationale": {"correct": "測試解析"},
                "sources": ["SYNTHETIC-TEST-SOURCE"],
            }
        ],
    }


def _run_cli(tmp_path, *, hide_answers=False, pair_mismatch=False):
    form_a = _exam("甲卷")
    form_b = _exam("乙卷")
    if pair_mismatch:
        form_b["items"][0]["pair_id"] = "P02"

    path_a = tmp_path / "form_a.json"
    path_b = tmp_path / "form_b.json"
    path_a.write_text(json.dumps(form_a, ensure_ascii=False), encoding="utf-8")
    path_b.write_text(json.dumps(form_b, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "docx"
    cmd = [
        sys.executable, "-m", "src.export_docx",
        str(path_a), str(path_b), "--out", str(out),
    ]
    if hide_answers:
        cmd.append("--hide-answers")
    result = subprocess.run(cmd, capture_output=True, text=True, check=False)
    return result, out


def test_cli_delivers_two_readable_word_files_with_answer_appendix(tmp_path):
    result, out = _run_cli(tmp_path)
    assert result.returncode == 0, result.stderr
    paths = sorted(out.glob("*.docx"))
    assert len(paths) == 2
    for path in paths:
        doc = Document(path)
        paragraphs = "\n".join(p.text for p in doc.paragraphs)
        assert "測試題幹" in paragraphs
        assert "答案與解析" in paragraphs
        assert "正解解析：測試解析" in paragraphs
        assert "來源 ID：SYNTHETIC-TEST-SOURCE" in paragraphs
        assert doc.tables[0].rows[1].cells[1].text == "C"


def test_cli_supports_answer_hidden_word_files(tmp_path):
    result, out = _run_cli(tmp_path, hide_answers=True)
    assert result.returncode == 0, result.stderr
    paths = sorted(out.glob("*.docx"))
    assert len(paths) == 2
    for path in paths:
        doc = Document(path)
        paragraphs = "\n".join(p.text for p in doc.paragraphs)
        assert "測試題幹" in paragraphs
        assert "答案與解析" not in paragraphs
        assert "正解解析" not in paragraphs
        assert not doc.tables


def test_cli_rejects_mismatched_pairs_without_creating_word_files(tmp_path):
    result, out = _run_cli(tmp_path, pair_mismatch=True)
    assert result.returncode != 0
    assert "Pair ID mismatch" in result.stderr
    assert not list(out.glob("*.docx"))
