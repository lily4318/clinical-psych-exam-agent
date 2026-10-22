"""Structural acceptance tests; fixtures are synthetic and never released as exam items."""

import json

import pytest
from docx import Document

from src.verify_external_delivery import verify_pair, main


def _make_form(path, form, n=2, *, bad_option=False, missing_source=False,
               answer_conflict=False, version="external-test-v0"):
    doc = Document()
    doc.add_paragraph(f"臨床心理師國考模擬測驗｜{form}")
    doc.add_paragraph(f"版本：{version} 題數：{n} 題")
    doc.add_paragraph("未經獨立專家驗證的測試文件")
    for k in range(1, n + 1):
        doc.add_paragraph(f"{k}. {form} 測試題目 {k}？")
        for letter in "ABCD":
            if bad_option and k == 1 and letter == "D":
                continue
            doc.add_paragraph(f"({letter}) 選項內容 {letter}")
    doc.add_paragraph(f"{form}｜答案與解析")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "題號"
    table.rows[0].cells[1].text = "答案"
    for k in range(1, n + 1):
        row = table.add_row()
        row.cells[0].text = str(k)
        row.cells[1].text = "B" if answer_conflict and k == 1 else "A"
        doc.add_paragraph(f"第 {k} 題｜答案 A")
        doc.add_paragraph("正解解析：僅供測試的結構化說明")
        if not (missing_source and k == 1):
            doc.add_paragraph("資料來源 ID：UNIT-TEST-ONLY")
    doc.save(path)
    return path


def _pair(tmp_path, **kwargs):
    a = _make_form(tmp_path / "a.docx", "甲卷", **kwargs)
    b = _make_form(tmp_path / "b.docx", "乙卷", **kwargs)
    return a, b


def test_external_style_pair_passes_with_unknown_links_explicit(tmp_path):
    result = verify_pair(*_pair(tmp_path), 2)
    assert result["ok"]
    assert result["files"]["form_a"]["questions"] == 2
    assert result["files"]["form_b"]["answers"] == 2
    assert result["host_download_links"] == "not_verified_by_local_file_check"
    assert result["clinical_correctness_and_source_support"] == "not_independently_verified"


def test_missing_file_fails_closed(tmp_path):
    a, b = _pair(tmp_path)
    b.unlink()
    result = verify_pair(a, b, 2)
    assert not result["ok"]
    assert result["files"]["form_b"]["errors"]


def test_invalid_docx_fails_closed(tmp_path):
    a, b = _pair(tmp_path)
    b.write_text("not a docx", encoding="utf-8")
    assert not verify_pair(a, b, 2)["ok"]


def test_requested_item_count_not_silently_ignored(tmp_path):
    assert not verify_pair(*_pair(tmp_path), 3)["ok"]


def test_missing_option_fails(tmp_path):
    assert not verify_pair(*_pair(tmp_path, bad_option=True), 2)["ok"]


def test_missing_per_item_source_fails(tmp_path):
    assert not verify_pair(*_pair(tmp_path, missing_source=True), 2)["ok"]


def test_answer_table_conflict_fails(tmp_path):
    assert not verify_pair(*_pair(tmp_path, answer_conflict=True), 2)["ok"]


def test_mismatched_versions_fail(tmp_path):
    a, b = _pair(tmp_path)
    _make_form(b, "乙卷", version="different-v1")
    result = verify_pair(a, b, 2)
    assert not result["ok"]
    assert "版本不同" in str(result["files"]["form_b"]["errors"])


def test_same_document_or_wrong_form_fails(tmp_path):
    a, b = _pair(tmp_path)
    assert not verify_pair(a, a, 2)["ok"]
    assert not verify_pair(b, a, 2)["ok"]


def test_same_parallel_stems_fail(tmp_path):
    a, b = _pair(tmp_path)
    _make_form(b, "乙卷")
    doc = Document(b)
    for p in doc.paragraphs:
        if p.text.startswith("1. 乙卷"):
            p.text = "1. 甲卷 測試題目 1？"
    doc.save(b)
    assert not verify_pair(a, b, 2)["ok"]


def test_invalid_item_count_fails(tmp_path):
    with pytest.raises(ValueError):
        verify_pair(*_pair(tmp_path), 0)


def test_cli_emits_machine_readable_status(tmp_path, monkeypatch, capsys):
    a, b = _pair(tmp_path)
    monkeypatch.setattr("sys.argv", ["verify_external_delivery", str(a), str(b), "--items", "2"])
    assert main() == 0
    record = json.loads(capsys.readouterr().out)
    assert record["ok"] is True
    monkeypatch.setattr("sys.argv", ["verify_external_delivery", str(a), str(b), "--items", "3"])
    assert main() == 1


def test_repository_exporter_is_accepted(tmp_path):
    """The portable verifier must accept the repo's native DOCX output too."""
    from src.export_docx import export_parallel_docx

    def exam(form):
        return {
            "form": form,
            "version": "unit-test-unverified",
            "approved_by_user": False,
            "quality_status": "未經獨立專家驗證之模擬草稿",
            "items": [
                {
                    "pair_id": f"P{i:02}",
                    "stem": f"{form} 測試題目 {i}",
                    "options": {letter: f"測試選項 {letter}" for letter in "ABCD"},
                    "correct_answer": "A",
                    "rationale": {
                        "correct": "這是測試解析。",
                        **{letter: f"{letter} 選項測試解析。" for letter in "ABCD"},
                    },
                    "sources": ["UNIT-TEST-ONLY"],
                }
                for i in (1, 2)
            ],
        }

    paths = export_parallel_docx(exam("甲卷"), exam("乙卷"), tmp_path)
    assert verify_pair(paths["form_a"], paths["form_b"], 2)["ok"]
