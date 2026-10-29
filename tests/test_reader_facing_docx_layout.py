"""Reader-facing paper layout; synthetic items are layout fixtures, not clinical exams."""

from docx import Document

from src.export_docx import export_exam_docx, export_parallel_docx
from src.verify_external_delivery import verify_pair


def _paper(form, count=3, *, offset=0):
    names = {"ADHD": "注意力不足／過動症", "ASD": "自閉症"}
    items = []
    for index in range(count):
        key = "ADHD" if (index + offset) % 2 == 0 else "ASD"
        items.append({
            "pair_id": f"P{index+1:02d}",
            "stem": f"{form} 格式測試題幹 {index+1}？",
            "options": {letter: f"測試選項 {letter}" for letter in "ABCD"},
            "correct_answer": "A",
            "rationale": {
                "correct": "這是排版測試用正解說明，不構成臨床主張。",
                "A": "A 選項排版測試。",
                "B": "B 選項排版測試。",
                "C": "C 選項排版測試。",
                "D": "D 選項排版測試。",
            },
            "sources": ["TEST-ADHD" if key == "ADHD" else "TEST-ASD"],
            "category": key,
        })
    return {
        "form": form,
        "version": "reader-layout-test-v1",
        "domain_id": "child_adolescent",
        "domain_name": "兒童／青少年臨床心理",
        "category_labels": names,
        "source_catalog": {
            "TEST-ADHD": {"title": "測試用參考資料（非臨床證據）", "url": "https://example.org/test-adhd"},
            "TEST-ASD": {"title": "第二份測試用參考資料（非臨床證據）", "url": "https://example.org/test-asd"},
            "NOT-CITED": {"title": "不應出現在試卷上的未引用資料", "url": "https://example.org/no"},
        },
        "quality_status": "未經獨立專家驗證之模擬草稿",
        "approved_by_user": False,
        "items": items,
    }


def test_clean_layout_and_appendix_order(tmp_path):
    doc_path = export_exam_docx(_paper("甲卷"), tmp_path)
    doc = Document(doc_path)
    texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    joined = "\n".join(texts)

    question_index = next(i for i, x in enumerate(texts) if "1. 甲卷" in x)
    scope_index = texts.index("命題範圍與依據")
    answers_index = texts.index("答案一覽")
    rationale_index = texts.index("逐題解析")
    source_index = texts.index("來源對照與參考資料")
    last_rationale_index = max(i for i, x in enumerate(texts) if x.startswith("正解解析："))

    assert question_index < scope_index < answers_index < rationale_index
    assert last_rationale_index < source_index
    assert "命題範圍：兒童／青少年臨床心理" in joined
    assert "考點配置：注意力不足／過動症 2 題、自閉症 1 題" in joined
    assert "第 1 題：TEST-ADHD" in joined
    assert "第 2 題：TEST-ASD" in joined
    assert "[TEST-ADHD] 測試用參考資料" in joined
    assert "https://example.org/test-asd" in joined
    assert "不應出現在試卷上的未引用資料" not in joined
    assert "非官方模擬試題" in joined
    assert "來源 ID：" not in joined  # moved out of individual explanations

    for not_reader_facing in (
        "版本：", "自動產生之模擬測驗", "未標示為使用者正式核可版本",
        "StubQuestionProvider", "system_prompt.md", "Agent 規則",
        "正式核可流程", "source_catalog",
    ):
        assert not_reader_facing not in joined
    assert doc.core_properties.identifier == "exam_version:reader-layout-test-v1"
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 4


def test_two_forms_with_distinct_allocations_do_not_share_scope_data(tmp_path):
    a = _paper("甲卷", 3, offset=0)
    b = _paper("乙卷", 3, offset=1)
    exported = export_parallel_docx(a, b, tmp_path)
    texts_a = "\n".join(x.text for x in Document(exported["form_a"]).paragraphs)
    texts_b = "\n".join(x.text for x in Document(exported["form_b"]).paragraphs)
    assert "注意力不足／過動症 2 題、自閉症 1 題" in texts_a
    assert "自閉症 2 題、注意力不足／過動症 1 題" in texts_b
    report = verify_pair(exported["form_a"], exported["form_b"], 3)
    assert report["ok"], report


def test_question_only_version_has_no_explanations_or_technical_metadata(tmp_path):
    doc = Document(export_exam_docx(_paper("甲卷"), tmp_path, include_answers=False))
    text = "\n".join(x.text for x in doc.paragraphs)
    assert "命題範圍與依據" not in text
    assert "答案與解析" not in text
    assert "來源對照與參考資料" not in text
    assert "版本：" not in text
    assert "甲卷 格式測試題幹" in text


def test_unknown_bibliographic_entries_are_not_invented(tmp_path):
    exam = _paper("甲卷", 1)
    exam.pop("source_catalog")
    text = "\n".join(p.text for p in Document(export_exam_docx(exam, tmp_path)).paragraphs)
    assert "第 1 題：TEST-ADHD" in text
    assert "測試用參考資料" not in text
    assert "https://example.org" not in text
