from __future__ import annotations

import argparse
import json
import os
import re
from collections import Counter
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .verify_external_delivery import normalize_question_stem


def _safe_slug(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "-", value).strip(".-") or "draft"


def _validate_exam(exam: dict[str, Any], *, include_answers: bool) -> None:
    if exam.get("form") not in {"甲卷", "乙卷"}:
        raise ValueError("Expected a 甲卷 or 乙卷 exam.")
    items = exam.get("items")
    if not isinstance(items, list) or not items:
        raise ValueError("Exam must contain a non-empty items list.")
    for number, item in enumerate(items, start=1):
        if not isinstance(item, dict) or not str(item.get("stem", "")).strip():
            raise ValueError(f"Item {number} requires a question stem.")
        options = item.get("options")
        if not isinstance(options, dict) or any(
            not str(options.get(letter, "")).strip() for letter in "ABCD"
        ):
            raise ValueError(f"Item {number} requires non-empty A-D options.")
        if include_answers:
            answer = item.get("correct_answer")
            if not isinstance(answer, str) or answer not in {"A", "B", "C", "D"}:
                raise ValueError(f"Item {number} requires one valid answer letter.")
            rationale = item.get("rationale")
            if not isinstance(rationale, dict):
                raise ValueError(f"Item {number} requires a correct-answer rationale.")
            for key in ("correct", *"ABCD"):
                detail = rationale.get(key)
                if not isinstance(detail, str) or not detail.strip():
                    raise ValueError(
                        f"Item {number} requires non-empty {key} rationale."
                    )
            sources = item.get("sources")
            if (
                not isinstance(sources, list)
                or not sources
                or any(not isinstance(source, str) or not source.strip() for source in sources)
                or len({source.strip() for source in sources}) != len(sources)
            ):
                raise ValueError(
                    f"Item {number} requires unique non-empty source IDs."
                )


def _unique_sources(exam: dict[str, Any]) -> list[str]:
    return list(dict.fromkeys(
        str(source_id)
        for item in exam["items"]
        for source_id in (item.get("sources") or [])
        if str(source_id).strip()
    ))


def _add_scope_and_basis(document: Document, exam: dict[str, Any]) -> None:
    """Only reader-relevant exam-specific scope and evidence, never agent instructions."""
    document.add_heading("命題範圍與依據", level=1)

    domain = exam.get("domain_name")
    if not domain and exam.get("domain_id") == "child_adolescent":
        domain = "兒童／青少年臨床心理"
    if isinstance(domain, str) and domain.strip():
        document.add_paragraph(f"命題範圍：{domain.strip()}")

    category_labels = exam.get("category_labels") or {}
    allocations = Counter(
        str(item["category"])
        for item in exam["items"]
        if item.get("category")
    )
    if allocations:
        description = "、".join(
            f"{category_labels.get(key, key)} {count} 題"
            for key, count in allocations.items()
        )
        document.add_paragraph(f"考點配置：{description}")

    scope_note = exam.get("scope_note")
    if isinstance(scope_note, str) and scope_note.strip():
        document.add_paragraph(scope_note.strip())

    years = exam.get("historical_years")
    if isinstance(years, list) and years and all(
        isinstance(year, int) and not isinstance(year, bool) for year in years
    ):
        years = sorted(set(years))
        document.add_paragraph(
            f"歷屆試題參考年度：{years[0]}–{years[-1]} 年"
        )

    basis = exam.get("basis_note")
    if isinstance(basis, str) and basis.strip():
        document.add_paragraph(f"命題依據：{basis.strip()}")
    elif _unique_sources(exam):
        document.add_paragraph(
            "命題依據：各題所列之參考來源，詳見卷末「來源對照與參考資料」。"
        )


def _add_source_appendix(document: Document, exam: dict[str, Any]) -> None:
    """Source mapping after all answer explanations; never fabricate bibliographies."""
    used_ids = _unique_sources(exam)
    if not used_ids:
        return
    document.add_heading("來源對照與參考資料", level=1)
    document.add_heading("題號與來源 ID 對照", level=2)
    for number, item in enumerate(exam["items"], start=1):
        ids = [
            str(source_id).strip()
            for source_id in (item.get("sources") or [])
            if str(source_id).strip()
        ]
        document.add_paragraph(
            f"第 {number} 題：{', '.join(ids) if ids else '未提供來源 ID'}"
        )

    catalog = exam.get("source_catalog") or {}
    referenced = [
        (source_id, catalog.get(source_id))
        for source_id in used_ids
        if isinstance(catalog.get(source_id), dict)
    ]
    if referenced:
        document.add_heading("參考來源", level=2)
        for source_id, source in referenced:
            title = str(source.get("title") or "").strip()
            url = str(source.get("url") or "").strip()
            if title:
                document.add_paragraph(f"[{source_id}] {title}")
            if url.startswith("https://"):
                document.add_paragraph(url)


def export_exam_docx(
    exam: dict[str, Any],
    out_dir: str | Path,
    *,
    include_answers: bool = True,
) -> Path:
    """Export a reader-facing Word: questions first, analysis, then source mapping."""
    _validate_exam(exam, include_answers=include_answers)
    destination = Path(out_dir)
    destination.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")

    title = document.add_heading(f"臨床心理師模擬測驗｜{exam['form']}", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version = str(exam.get("version") or "draft")
    # Provenance remains machine-readable but is not printed as student-facing text.
    document.core_properties.identifier = f"exam_version:{version}"

    for number, item in enumerate(exam["items"], start=1):
        document.add_heading(f"{number}. {item['stem']}", level=2)
        for letter in "ABCD":
            document.add_paragraph(f"({letter}) {item['options'][letter]}")

    if include_answers:
        document.add_page_break()
        document.add_heading(f"{exam['form']}｜答案與解析", 0)
        _add_scope_and_basis(document, exam)
        document.add_heading("答案一覽", level=1)

        answer_table = document.add_table(rows=1, cols=2)
        answer_table.style = "Table Grid"
        answer_table.rows[0].cells[0].text = "題號"
        answer_table.rows[0].cells[1].text = "答案"
        for number, item in enumerate(exam["items"], start=1):
            cells = answer_table.add_row().cells
            cells[0].text = str(number)
            cells[1].text = item["correct_answer"]

        document.add_heading("逐題解析", level=1)
        for number, item in enumerate(exam["items"], start=1):
            document.add_heading(f"第 {number} 題｜答案 {item['correct_answer']}", level=2)
            rationale = item["rationale"]
            document.add_paragraph(f"正解解析：{rationale['correct']}")
            for letter in "ABCD":
                # Preflight guarantees all four explanations exist; never silently omit one.
                document.add_paragraph(f"{letter} 選項：{rationale[letter].strip()}")

        _add_source_appendix(document, exam)
        if exam.get("approved_by_user") is not True:
            document.add_paragraph(
                "本卷為非官方模擬試題；內容尚未經獨立臨床專家逐題驗證。"
            )

    slug = "form_a" if exam["form"] == "甲卷" else "form_b"
    name = f"{slug}_{_safe_slug(version)}.docx"
    path = destination / name
    document.save(path)
    return path


def export_parallel_docx(
    form_a: dict[str, Any],
    form_b: dict[str, Any],
    out_dir: str | Path,
    *,
    include_answers: bool = True,
) -> dict[str, Path]:
    """Export exactly two Word files; no formal-release approval is implied."""
    _validate_exam(form_a, include_answers=include_answers)
    _validate_exam(form_b, include_answers=include_answers)
    if form_a["form"] != "甲卷" or form_b["form"] != "乙卷":
        raise ValueError("Expected form_a=甲卷 and form_b=乙卷.")
    if len(form_a["items"]) != len(form_b["items"]):
        raise ValueError("Parallel forms must have the same item count.")
    if str(form_a.get("version") or "draft") != str(form_b.get("version") or "draft"):
        raise ValueError("Parallel forms must use the same version.")
    seen_pair_ids: set[str] = set()
    for number, (item_a, item_b) in enumerate(
        zip(form_a["items"], form_b["items"]), start=1
    ):
        pair_a, pair_b = item_a.get("pair_id"), item_b.get("pair_id")
        if (
            not isinstance(pair_a, str)
            or not pair_a.strip()
            or not isinstance(pair_b, str)
            or not pair_b.strip()
        ):
            raise ValueError(f"Missing pair ID at item {number}.")
        if pair_a != pair_b:
            raise ValueError(f"Pair ID mismatch at item {number}.")
        if pair_a in seen_pair_ids:
            raise ValueError(f"Duplicate pair ID at item {number}: {pair_a}.")
        seen_pair_ids.add(pair_a)

    # The JSON-to-DOCX CLI must reject reused stems before writing either
    # file. This mirrors the external G3 verifier's exact-text normalization
    # (NFKC, whitespace, case); it is not a semantic-equivalence check.
    seen_stems: dict[str, tuple[str, int]] = {}
    for label, exam in (("甲卷", form_a), ("乙卷", form_b)):
        for number, item in enumerate(exam["items"], start=1):
            key = normalize_question_stem(item["stem"])
            if key in seen_stems:
                prior_label, prior_number = seen_stems[key]
                raise ValueError(
                    "Duplicate question stem in parallel forms: "
                    f"{prior_label} item {prior_number}, {label} item {number}."
                )
            seen_stems[key] = (label, number)

    # The public JSON-to-DOCX CLI must not leave only one Word if the second
    # export fails, nor silently replace a previously delivered exam. Stage
    # both files on the same filesystem, then publish each with an exclusive
    # hard link (os.link fails if the destination already exists).
    destination = Path(out_dir)
    destination.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix=".parallel-docx-stage-", dir=destination) as temp:
        staged = {
            "form_a": export_exam_docx(form_a, temp, include_answers=include_answers),
            "form_b": export_exam_docx(form_b, temp, include_answers=include_answers),
        }
        final = {key: destination / path.name for key, path in staged.items()}
        if final["form_a"] == final["form_b"]:
            raise ValueError("Parallel Word output filenames must be distinct.")
        for path in final.values():
            if path.exists():
                raise FileExistsError(f"Refusing to overwrite an existing Word: {path}")

        published: list[Path] = []
        try:
            for key in ("form_a", "form_b"):
                # Exclusive creation avoids a race between exists() and replace().
                # Hard-linking staged bytes is atomic on the same filesystem.
                os.link(staged[key], final[key])
                published.append(final[key])
        except Exception:
            for path in published:
                path.unlink(missing_ok=True)
            raise
        return final


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Export an existing validated parallel pair to two Word documents."
    )
    parser.add_argument("form_a", type=Path, help="甲卷 JSON")
    parser.add_argument("form_b", type=Path, help="乙卷 JSON")
    parser.add_argument("--out", type=Path, default=Path("output/docx"))
    parser.add_argument("--hide-answers", action="store_true")
    args = parser.parse_args()
    form_a = json.loads(args.form_a.read_text(encoding="utf-8"))
    form_b = json.loads(args.form_b.read_text(encoding="utf-8"))
    paths = export_parallel_docx(
        form_a, form_b, args.out, include_answers=not args.hide_answers
    )
    for name, path in paths.items():
        print(f"{name}: {path}")


if __name__ == "__main__":
    main()
