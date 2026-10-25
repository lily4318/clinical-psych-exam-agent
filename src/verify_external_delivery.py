"""Provider-neutral structural checks for externally generated parallel DOCX files.

This validates saved artifacts, NOT their download links or clinical correctness.
No model inference, external network requests, or paid services are used.
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path
from typing import Any

from docx import Document

QUESTION = re.compile(r"^\s*(\d+)[.．、]\s*(\S.*)$")
OPTION = re.compile(r"^\s*[(（]([ABCD])[)）]\s*(\S.*)$", re.I)
EXPLAINED_ANSWER = re.compile(r"^\s*第\s*(\d+)\s*題\s*[｜|：:]\s*答案\s*([ABCD])\s*$")
SHORT_ANSWER = re.compile(r"^\s*(\d+)[.．]\s*([ABCD])\s*$")
VERSION = re.compile(r"版本[：:]\s*([^\s　]+)")
SOURCE_MAPPING = re.compile(r"^第\s*(\d+)\s*題\s*[：:]\s*(\S.*)$")
ANSWER_HEADING = re.compile(r"^[甲乙]卷\s*[｜|]\s*答案與解析")
OPTION_RATIONALE = re.compile(r"^\s*([ABCD])(?:\s*選項)?\s*[：:]\s*(.*)$", re.I)
READER_SECTIONS = ("命題範圍與依據", "答案一覽", "逐題解析", "來源對照與參考資料")


def normalize_question_stem(stem: str) -> str:
    """Normalize only exact-text variants; not a semantic similarity judgement."""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", stem)).casefold()


def _parse_answers(doc: Any, after: list[str], expected_count: int) -> tuple[dict[int, str], list[str]]:
    errors: list[str] = []
    from_headings: dict[int, str] = {}
    for text in after:
        match = EXPLAINED_ANSWER.match(text) or SHORT_ANSWER.match(text)
        if match:
            n, letter = int(match.group(1)), match.group(2).upper()
            if n in from_headings:
                errors.append(f"重複的答案解析標題：{n}")
            from_headings[n] = letter

    from_table: dict[int, str] = {}
    matching_tables = 0
    for table in doc.tables:
        if len(table.columns) < 2 or len(table.rows) == 0:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells[:2]]
        if "題號" not in headers[0] or "答案" not in headers[1]:
            continue
        matching_tables += 1
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells[:2]]
            if not cells[0].isdigit() or cells[1].upper() not in "ABCD" or len(cells[1]) != 1:
                errors.append(f"答案表有無效題號／選項：{cells}")
                continue
            n = int(cells[0])
            if n in from_table:
                errors.append(f"答案表有重複題號：{n}")
            from_table[n] = cells[1].upper()
    if matching_tables > 1:
        errors.append("偵測到多個答案表，無法確認唯一答案")
    if from_table and from_headings and from_table != from_headings:
        errors.append("答案表與逐題解析標題不一致")
    answers = from_table or from_headings
    if set(answers) != set(range(1, expected_count + 1)):
        errors.append("答案題號不完整或與指定題數不符")
    return answers, errors



def _check_modern_reader_layout(
    doc: Any, questions: list[str], answers: list[str], expected_count: int
) -> tuple[list[str], int]:
    """Validate new reader-facing layout; not a clinical or citation-content audit."""
    errors: list[str] = []
    positions: dict[str, int] = {}
    for section in READER_SECTIONS:
        indices = [i for i, line in enumerate(answers) if line == section]
        if len(indices) != 1:
            errors.append(f"新版考卷缺少或重複章節：{section}")
        else:
            positions[section] = indices[0]
    if len(positions) == len(READER_SECTIONS):
        indices = [positions[section] for section in READER_SECTIONS]
        if indices != sorted(indices):
            errors.append("新版考卷章節順序錯誤：命題範圍→答案一覽→逐題解析→來源對照")
        if positions["命題範圍與依據"] < positions["答案一覽"]:
            if not any(
                answers[i].startswith(("命題範圍：", "考點配置：", "命題依據："))
                for i in range(positions["命題範圍與依據"] + 1, positions["答案一覽"])
            ):
                errors.append("命題範圍與依據區沒有本卷實際內容")

    if any(
        line.startswith(("版本：", "未標示為使用者正式核可版本", "自動產生之模擬測驗"))
        for line in questions
    ):
        errors.append("試題區含有內部版本或 AI 生成說明")

    if not any(
        len(table.rows) >= expected_count + 1
        and len(table.columns) >= 2
        and table.rows[0].cells[0].text.strip() == "題號"
        and table.rows[0].cells[1].text.strip() == "答案"
        for table in doc.tables
    ):
        errors.append("新版考卷缺少完整答案一覽表")

    start = positions.get("逐題解析", -1)
    stop = positions.get("來源對照與參考資料", len(answers))
    if start < 0:
        return errors, 0
    if stop <= start:
        errors.append("來源對照必須放在全部逐題解析之後")
        return errors, 0

    headers: list[tuple[int, int]] = []
    for i in range(start + 1, stop):
        match = EXPLAINED_ANSWER.match(answers[i]) or SHORT_ANSWER.match(answers[i])
        if match:
            headers.append((i, int(match.group(1))))
    if [n for _, n in headers] != list(range(1, expected_count + 1)):
        errors.append("新版逐題解析的題號不完整、重複或順序錯誤")

    completed = 0
    for index, (position, number) in enumerate(headers):
        end = headers[index + 1][0] if index + 1 < len(headers) else stop
        chunk = answers[position + 1:end]
        correct = [line for line in chunk if line.startswith("正解解析：") and line.split("：", 1)[-1].strip()]
        if len(correct) != 1:
            errors.append(f"第 {number} 題正解解析缺漏或重複")
        explanations: dict[str, str] = {}
        for line in chunk:
            match = OPTION_RATIONALE.match(line)
            if match:
                letter, detail = match.group(1).upper(), match.group(2).strip()
                if letter in explanations:
                    errors.append(f"第 {number} 題 {letter} 選項解析重複")
                explanations[letter] = detail
        missing = [key for key in "ABCD" if not explanations.get(key)]
        if missing:
            errors.append(f"第 {number} 題缺少 A–D 逐項解析：{','.join(missing)}")
        if len(correct) == 1 and not missing:
            completed += 1
    return errors, completed


def inspect_form(path: str | Path, form: str, item_count: int) -> dict[str, Any]:
    """Check a DOCX produced by either this repo or an external AI."""
    path = Path(path)
    errors: list[str] = []
    result: dict[str, Any] = {"form": form, "file": path.name, "ok": False,
                               "questions": 0, "answers": 0, "version": None, "errors": errors,
                               "reader_layout": "unknown", "complete_option_rationales": None,
                               "warnings": []}
    if path.suffix.lower() != ".docx" or not path.is_file() or path.stat().st_size == 0:
        errors.append("檔案不存在、為空或不是 DOCX")
        return result
    try:
        doc = Document(path)
    except Exception as exc:
        errors.append(f"DOCX 無法開啟：{type(exc).__name__}")
        return result

    lines = [p.text.strip() for p in doc.paragraphs]
    if not any(form in text and "模擬" in text for text in lines[:6]):
        errors.append("文件開頭缺少正確卷別的模擬考標題")
    early_text = "\n".join(lines[:10])
    match = VERSION.search(early_text)
    identifier = doc.core_properties.identifier or ""
    hidden_version = identifier.removeprefix("exam_version:") if identifier.startswith("exam_version:") else None
    result["version"] = match.group(1) if match else hidden_version
    # New reader-facing papers put version in DOCX properties, not the question text.
    # External documents without version evidence remain distinguishable as unknown.
    if not ("未經獨立專家驗證" in "\n".join(lines) or
            "未標示為使用者正式核可版本" in "\n".join(lines) or
            "未經獨立臨床專家逐題驗證" in "\n".join(lines)):
        errors.append("缺少未經獨立專家驗證／非正式核可標示")

    answer_positions = [i for i, text in enumerate(lines) if ANSWER_HEADING.match(text)]
    if len(answer_positions) != 1 or form not in lines[answer_positions[0]]:
        errors.append("缺少正確卷別的單一『答案與解析』區段")
        return result
    split = answer_positions[0]
    question_lines, answer_lines = lines[:split], lines[split + 1:]
    modern = identifier.startswith("exam_version:") or any(
        section in answer_lines for section in READER_SECTIONS
    )
    result["reader_layout"] = "modern_strict" if modern else "legacy_compatible"
    if modern:
        layout_errors, completed = _check_modern_reader_layout(
            doc, question_lines, answer_lines, item_count
        )
        errors.extend(layout_errors)
        result["complete_option_rationales"] = completed
    else:
        result["warnings"].append(
            "舊版文件僅通過相容性結構檢查；未驗證 A–D 逐項解析及新版章節順序"
        )
    question_stems: list[str] = []
    question_positions = []
    for i, text in enumerate(question_lines):
        m = QUESTION.match(text)
        if m and not OPTION.match(text):
            question_positions.append((i, int(m.group(1))))
            question_stems.append(m.group(2))
    nums = [number for _, number in question_positions]
    result["questions"] = len(nums)
    result["stems"] = question_stems
    if nums != list(range(1, item_count + 1)):
        errors.append(f"題號／數量不符：應有 1–{item_count}，實際 {nums}")
    for pos, (i, number) in enumerate(question_positions):
        end = question_positions[pos + 1][0] if pos + 1 < len(question_positions) else len(question_lines)
        options = [OPTION.match(s).group(1).upper() for s in question_lines[i + 1:end] if OPTION.match(s)]
        if sorted(options) != list("ABCD"):
            errors.append(f"第 {number} 題缺少或重複 A–D 選項：{options}")

    answers, answer_errors = _parse_answers(doc, answer_lines, item_count)
    result["answers"] = len(answers)
    errors.extend(answer_errors)
    explanation_lines = [line for line in answer_lines if line.startswith("正解解析：") and line.split("：", 1)[-1].strip()]
    if len(explanation_lines) != item_count:
        errors.append(f"有效正解解析數量不符：{len(explanation_lines)}/{item_count}")
    source_lines = [line for line in answer_lines if (line.startswith("來源 ID：") or line.startswith("資料來源 ID：")) and line.split("：", 1)[-1].strip()]
    appendix_positions = [i for i, line in enumerate(answer_lines) if line == "來源對照與參考資料"]
    appendix_sources = {}
    if len(appendix_positions) == 1:
        appendix = answer_lines[appendix_positions[0] + 1:]
        for line in appendix:
            match = SOURCE_MAPPING.match(line)
            if match:
                number, refs = int(match.group(1)), match.group(2).strip()
                if number in appendix_sources:
                    errors.append(f"來源對照重複題號：{number}")
                appendix_sources[number] = refs
    if source_lines:
        if len(source_lines) != item_count:
            errors.append(f"逐題來源 ID 數量不符：{len(source_lines)}/{item_count}")
    elif set(appendix_sources) != set(range(1, item_count + 1)):
        errors.append("卷末題號與來源 ID 對照不完整")
    if appendix_sources and any("未提供來源 ID" in refs for refs in appendix_sources.values()):
        errors.append("卷末來源對照不可含未提供來源 ID")
    if modern:
        if source_lines:
            errors.append("新版 Word 應將來源 ID 集中於全部解析後的來源對照")
        if set(appendix_sources) != set(range(1, item_count + 1)):
            errors.append("新版卷末來源對照題號不完整")
    result["ok"] = not errors
    return result


def verify_pair(form_a: str | Path, form_b: str | Path, item_count: int) -> dict[str, Any]:
    if not isinstance(item_count, int) or isinstance(item_count, bool) or item_count < 1:
        raise ValueError("--items must be a positive integer")
    a = inspect_form(form_a, "甲卷", item_count)
    b = inspect_form(form_b, "乙卷", item_count)
    if Path(form_a).resolve() == Path(form_b).resolve():
        b["errors"].append("甲乙卷不可使用同一檔案")
    if a["version"] and b["version"] and a["version"] != b["version"]:
        b["errors"].append("甲乙版本不同")
    # Check ALL question positions across both forms, not only aligned A/B pairs.
    # This detects copy-pasted items in the same form and cross-form reuse.
    seen_stems: dict[str, tuple[str, int]] = {}
    for result, label in ((a, "甲卷"), (b, "乙卷")):
        for number, stem in enumerate(result.get("stems", []), start=1):
            normalized = normalize_question_stem(stem)
            if not normalized:
                continue
            if normalized in seen_stems:
                previous_form, previous_number = seen_stems[normalized]
                result["errors"].append(
                    f"{label}第 {number} 題與{previous_form}第 {previous_number} 題"
                    "題幹重複（忽略空白及全半形差異）"
                )
            else:
                seen_stems[normalized] = (label, number)
    for doc in (a, b):
        doc.pop("stems", None)
        doc["ok"] = not doc["errors"]
    return {"ok": a["ok"] and b["ok"], "items_per_form": item_count,
            "files": {"form_a": a, "form_b": b},
            "host_download_links": "not_verified_by_local_file_check",
            "clinical_correctness_and_source_support": "not_independently_verified",
            "parallel_psychometric_equivalence": "not_verified"}


def main() -> int:
    parser = argparse.ArgumentParser(description="Check two externally generated DOCX files without any model API.")
    parser.add_argument("form_a", type=Path)
    parser.add_argument("form_b", type=Path)
    parser.add_argument("--items", type=int, required=True, help="Expected number of questions in EACH form")
    parser.add_argument("--output", type=Path, help="Optional machine-readable JSON report path")
    args = parser.parse_args()
    try:
        result = verify_pair(args.form_a, args.form_b, args.items)
    except ValueError as exc:
        parser.error(str(exc))
    output = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(output + "\n", encoding="utf-8")
    print(output)
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
